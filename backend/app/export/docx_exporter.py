"""DOCX Exporter for generated school schedules.

Generates a Microsoft Word (.docx) document matching the school's master schedule format:
- Landscape layout (A4 / standard landscape with narrow margins).
- Header:
  - Row 1: 'Предмет' (rowspan=2) | 'ФИО' (rowspan=2) | Day names ('Понедельник' .. 'Пятница', colspan=9 each).
  - Row 2: Period numbers 1-9 repeated under each day.
- Data Rows:
  - One row per (Subject, Teacher) pair.
  - Cells at (Day, Period) filled with formatted class label (e.g. '6Е' or '6Е-1' for subgroups).
"""

import logging
from collections import defaultdict
from typing import Dict, List, Literal, Optional, Set, Tuple

import docx
from docx.enum.section import WD_ORIENT, WD_SECTION_START
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor

from app.models.schema import ScheduleSlot, Teacher

logger = logging.getLogger(__name__)

DAYS: List[Literal["Понедельник", "Вторник", "Среда", "Четверг", "Пятница"]] = [
    "Понедельник",
    "Вторник",
    "Среда",
    "Четверг",
    "Пятница",
]
PERIODS: List[int] = list(range(1, 10))  # 1 to 9 (45 period columns)


def _set_cell_shading(cell, color_hex: str):
    """Set background color for a table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    tc_pr.append(shd)


def _set_cell_margins(cell, top=60, bottom=60, left=80, right=80):
    """Set inner cell padding in twips (1 pt = 20 twips)."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tc_pr.append(tc_mar)


def _set_table_borders(table):
    """Apply thin borders to the whole table."""
    tbl_pr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="D1D5DB"/>'
        f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="D1D5DB"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="E5E7EB"/>'
        f'<w:insideV w:val="single" w:sz="4" w:space="0" w:color="E5E7EB"/>'
        f'</w:tblBorders>'
    )
    tbl_pr.append(borders)


def format_slot_text(slot: ScheduleSlot) -> str:
    """Format slot for timetable display: e.g. '6Е' or '6Е-1' if subgroup is present."""
    if slot.group:
        return f"{slot.class_name}-{slot.group}"
    return slot.class_name


def export_schedule_to_docx(
    schedule: List[ScheduleSlot],
    teachers: List[Teacher],
    output_path: str,
) -> None:
    """Exports the schedule into a formatted DOCX table matching the school's master schedule format.

    Args:
        schedule: List of scheduled lesson slots.
        teachers: List of Teacher domain models for resolving full names.
        output_path: File path to write the generated .docx.
    """
    doc = docx.Document()

    # Step 1: Configure Page Setup to Landscape with narrow margins
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11.69)   # A4 landscape width (297mm)
    section.page_height = Inches(8.27)   # A4 landscape height (210mm)
    section.top_margin = Inches(0.4)
    section.bottom_margin = Inches(0.4)
    section.left_margin = Inches(0.4)
    section.right_margin = Inches(0.4)

    # Document Header Title
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_after = Pt(6)
    title_run = title_p.add_run("Расписание уроков")
    title_run.font.name = "Times New Roman"
    title_run.font.size = Pt(14)
    title_run.font.bold = True
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Step 2: Build Teacher ID -> Full Name map & group schedule slots
    teacher_name_map: Dict[str, str] = {t.id: t.full_name for t in teachers}

    # Group scheduled slots by (subject, teacher_id)
    # Mapping: (subject, teacher_id) -> (day, period) -> list[formatted_class_strings]
    row_grid: Dict[Tuple[str, str], Dict[Tuple[str, int], List[str]]] = defaultdict(
        lambda: defaultdict(list)
    )
    for slot in schedule:
        cell_text = format_slot_text(slot)
        row_grid[(slot.subject, slot.teacher_id)][(slot.day, slot.period)].append(cell_text)

    # Sort rows nicely: by subject name first, then by teacher's full name
    sorted_row_keys = sorted(
        row_grid.keys(),
        key=lambda k: (k[0].lower(), teacher_name_map.get(k[1], k[1]).lower()),
    )

    # Total table columns: Col 0 (Предмет) + Col 1 (ФИО) + 5 days * 9 periods = 47 columns
    total_cols = 2 + len(DAYS) * len(PERIODS)
    total_rows = 2 + len(sorted_row_keys)

    table = doc.add_table(rows=total_rows, cols=total_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_borders(table)

    # Force fixed column widths: without this, Word ignores per-cell widths and
    # autofits every column to an equal share of the table width, which is what
    # was silently defeating our column-width settings below.
    table.autofit = False

    # Proportional column widths matching existing_schedule.docx:
    # Col 0 (Subject): 1.0 in - wide enough for horizontal word-wrap over 2-3 lines
    # Col 1 (Teacher): 1.15 in (~1650 dxa) - compact teacher names with wrapping
    # Cols 2..46 (45 Period columns): 0.21 in (~302 dxa) each
    col_width_subj = Inches(1.0)
    col_width_teacher = Inches(1.15)
    col_width_period = Inches(0.21)

    table.columns[0].width = col_width_subj
    table.columns[1].width = col_width_teacher
    for c in range(2, total_cols):
        table.columns[c].width = col_width_period

    # Header styling constants
    header_bg_color = "F3F4F6"
    day_bg_colors = ["E0F2FE", "E0E7FF", "EDE9FE", "FAE8FF", "FCE7F3"]

    # Step 3: Populate Header Row 0
    # Col 0: "Предмет"
    cell_subj_hdr = table.cell(0, 0)
    cell_subj_hdr.text = "Предмет"
    _set_cell_shading(cell_subj_hdr, header_bg_color)
    cell_subj_hdr.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    # Col 1: "ФИО"
    cell_fio_hdr = table.cell(0, 1)
    cell_fio_hdr.text = "ФИО"
    _set_cell_shading(cell_fio_hdr, header_bg_color)
    cell_fio_hdr.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    # Days header (spanning 9 period columns each)
    for d_idx, day_name in enumerate(DAYS):
        start_col = 2 + d_idx * len(PERIODS)
        end_col = start_col + len(PERIODS) - 1
        day_color = day_bg_colors[d_idx % len(day_bg_colors)]

        # Set text on first cell of day block
        day_cell = table.cell(0, start_col)
        day_cell.text = day_name
        day_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

        # Shade all cells in the day span before merging
        for c_idx in range(start_col, end_col + 1):
            _set_cell_shading(table.cell(0, c_idx), day_color)

        # Merge cells horizontally across the 9 periods for this day
        if end_col > start_col:
            day_cell.merge(table.cell(0, end_col))

    # Step 4: Populate Header Row 1 (Period numbers 1-9 under each day)
    # Merge row 0 and row 1 vertically for 'Предмет' and 'ФИО'
    table.cell(0, 0).merge(table.cell(1, 0))
    table.cell(0, 1).merge(table.cell(1, 1))

    for d_idx in range(len(DAYS)):
        start_col = 2 + d_idx * len(PERIODS)
        day_color = day_bg_colors[d_idx % len(day_bg_colors)]
        for p_idx, period_num in enumerate(PERIODS):
            col = start_col + p_idx
            p_cell = table.cell(1, col)
            p_cell.text = str(period_num)
            _set_cell_shading(p_cell, day_color)
            p_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    # Format header text runs (bold, centered, Times New Roman 8pt)
    for r_idx in [0, 1]:
        for cell in table.rows[r_idx].cells:
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(8)
                    run.font.bold = True

    # Step 5: Populate Data Rows
    # Precompute (day, period) -> column index mapping
    slot_to_col: Dict[Tuple[str, int], int] = {
        (day, period): 2 + d_idx * len(PERIODS) + (period - 1)
        for d_idx, day in enumerate(DAYS)
        for period in PERIODS
    }

    for row_idx, (subject_name, teacher_id) in enumerate(sorted_row_keys, start=2):
        t_row = table.rows[row_idx]
        cells = t_row.cells

        # Set widths for the row
        cells[0].width = col_width_subj
        cells[1].width = col_width_teacher
        for c in range(2, total_cols):
            cells[c].width = col_width_period

        # Col 0: Subject - only set text on the first row of each consecutive subject run
        # so that subsequent vertical merging does not stack duplicate text runs.
        is_first_in_run = (row_idx == 2) or (sorted_row_keys[row_idx - 3][0] != subject_name)
        c_subj = cells[0]
        if is_first_in_run:
            c_subj.text = subject_name
            p0 = c_subj.paragraphs[0]
            p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p0.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(8)
                run.font.bold = True
        else:
            c_subj.text = ""

        c_subj.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        _set_cell_margins(c_subj, top=20, bottom=20, left=20, right=20)

        # Col 1: Teacher Full Name
        teacher_full_name = teacher_name_map.get(teacher_id, teacher_id)
        c_teacher = cells[1]
        c_teacher.text = teacher_full_name
        c_teacher.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        _set_cell_margins(c_teacher, top=20, bottom=20, left=30, right=30)
        p1 = c_teacher.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in p1.runs:
            run.font.name = "Times New Roman"
            run.font.size = Pt(7.5)

        # Populate scheduled slots for this (subject, teacher)
        grid_data = row_grid.get((subject_name, teacher_id), {})
        for (day, period), classes in grid_data.items():
            col = slot_to_col.get((day, period))
            if col is not None:
                cell = cells[col]
                cell.text = ", ".join(classes)
                _set_cell_shading(cell, "EFF6FF")
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                _set_cell_margins(cell, top=20, bottom=20, left=10, right=10)
                pp = cell.paragraphs[0]
                pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in pp.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(7)
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    # Step 6: Vertically merge Column 0 for consecutive rows of the same subject
    # and clean up extra empty paragraphs introduced by python-docx merge.
    curr_subj = None
    run_start_row = 2
    for r_idx in range(2, total_rows):
        subj = sorted_row_keys[r_idx - 2][0]
        if subj != curr_subj:
            if curr_subj is not None and (r_idx - 1) > run_start_row:
                merged_cell = table.cell(run_start_row, 0).merge(table.cell(r_idx - 1, 0))
                # Remove extra trailing empty paragraphs so text is not repeated or spaced out
                for p in merged_cell.paragraphs[1:]:
                    p._p.getparent().remove(p._p)
            curr_subj = subj
            run_start_row = r_idx
    if curr_subj is not None and (total_rows - 1) > run_start_row:
        merged_cell = table.cell(run_start_row, 0).merge(table.cell(total_rows - 1, 0))
        for p in merged_cell.paragraphs[1:]:
            p._p.getparent().remove(p._p)

    # Enable header repetition across pages
    tr_pr_0 = table.rows[0]._tr.get_or_add_trPr()
    tr_pr_0.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    tr_pr_1 = table.rows[1]._tr.get_or_add_trPr()
    tr_pr_1.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

    # Save document
    doc.save(output_path)
    logger.info("Schedule successfully exported to '%s' with %d data rows.", output_path, len(sorted_row_keys))
