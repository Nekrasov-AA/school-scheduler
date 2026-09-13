"""Curriculum parser for Russian school curriculum plans (Учебный план, .docx).

Converts УП_*.docx files into a list of CurriculumRequirement objects.
Supports three structural shapes corresponding to levels of Russian general education:
- NOO (НОО - Начальное общее образование, Grades 1-4)
- OOO (ООО - Основное общее образование, Grades 5-9)
- SOO (СОО - Среднее общее образование, Grades 10-11)
"""

import logging
import re
from typing import Any, Dict, List, Optional, Set, Tuple

import docx
from docx.table import Table

from app.models.schema import CurriculumRequirement

logger = logging.getLogger(__name__)

# Prefixes for rows that represent metadata, totals, or section banners (not actual subjects)
NON_SUBJECT_PREFIXES = (
    "итого",
    "максимальная",
    "количество",
    "обязательная часть",
    "часть, формируемая",
    "предметные области",
    "учебные предметы",
    "формы промежуточной",
    "всего",
)


def normalize_text(text: Optional[str]) -> str:
    """Strip extra whitespace, replace non-breaking spaces, and keep original Cyrillic casing."""
    if not text:
        return ""
    text = text.replace("\xa0", " ").replace("\n", " ")
    return re.sub(r"\s+", " ", text).strip()


def parse_hours(val: Optional[str]) -> Optional[float]:
    """Parse academic hours string into float.

    Handles standard integers, decimal commas/dots (0.5, 1,5), and in-depth notations ('5 угл', '3угл').
    Returns None if empty, non-numeric, or dash.
    """
    cleaned = normalize_text(val)
    if not cleaned or cleaned == "-":
        return None

    # Matches numbers like '5', '5.0', '1,5', '0.5', '5 угл', '3угл'
    m = re.search(r"^(\d+(?:[.,]\d+)?)(?:\s*угл)?", cleaned, re.IGNORECASE)
    if m:
        num_str = m.group(1).replace(",", ".")
        try:
            return float(num_str)
        except ValueError:
            return None
    return None


def is_non_subject_row(text: str) -> bool:
    """Check whether a text represents a header banner, section division, or summary row."""
    low = text.lower()
    if not low:
        return True
    if any(low.startswith(p) for p in NON_SUBJECT_PREFIXES):
        return True
    if any(k in low for k in ["2025-2026", "2024-2025", "2026-2027", "2027-2028", "формы промежуточной аттестации"]):
        return True
    return False


def detect_curriculum_shape(table: Table) -> str:
    """Detect which of the three curriculum table shapes (NOO/OOO/SOO) is present.

    DETECTION HEURISTIC:
    The heuristic examines table header rows (specifically row index 2 and surrounding header rows)
    which define the grade levels for the columns across Russian school curricula:
    1. Scan Row 2 (and rows 0..2) for grade numbers:
       - If grades present are in range [1, 4] and max grade <= 4: -> 'NOO' (Primary, 1-4)
       - If grades present are in range [5, 9] (or verticals exist): -> 'OOO' (Basic general, 5-9)
       - If grades present are in range [10, 11] (or high-school tracks/class letters): -> 'SOO' (Secondary, 10-11)
    2. Fallback checks based on column count and keywords:
       - 7 columns, single column per grade -> 'NOO'
       - 15+ columns, presence of 'вертикаль' -> 'OOO'
       - Presence of 'гуманитарный'/'технологический' profile tracks or class letters '11а'/'11б' -> 'SOO'
    """
    found_grades: Set[int] = set()
    has_verticals = False
    has_profiles = False

    # Check rows 0..2 for header indicators
    header_rows = min(len(table.rows), 3)
    for r_idx in range(header_rows):
        row = table.rows[r_idx]
        for cell in row.cells:
            txt = normalize_text(cell.text)
            if not txt:
                continue

            low = txt.lower()
            if "вертикаль" in low:
                has_verticals = True
            if any(p in low for p in ["гуманитарный", "социально-экономический", "технологический", "естественно-научный"]):
                has_profiles = True

            # Extract standalone grade numbers or numbers before year/class letter (e.g. '5 2025-2026', '11 б', '1')
            m_grades = re.findall(r"\b(1[01]|[1-9])(?:\s*(?:[а-яёa-z]|\d{4}-\d{4}))?\b", txt, re.IGNORECASE)
            for mg in m_grades:
                found_grades.add(int(mg))

    logger.debug("Shape detection - Grades found: %s, Verticals: %s, Profiles: %s, Cols: %d",
                 found_grades, has_verticals, has_profiles, len(table.columns))

    if found_grades:
        if max(found_grades) <= 4:
            return "NOO"
        if min(found_grades) >= 10 or has_profiles:
            return "SOO"
        if max(found_grades) <= 9 or has_verticals:
            return "OOO"

    # Secondary heuristic based on column count and keywords
    if len(table.columns) >= 15 or has_verticals:
        return "OOO"
    if has_profiles or len(table.columns) == 13:
        return "SOO"

    return "NOO"


def _parse_noo_table(table: Table) -> List[CurriculumRequirement]:
    """Parse Primary General Education (НОО, Grades 1-4) curriculum table.

    Layout:
    Col 0: Subject Area (Предметные области)
    Col 1: Subject Name (Учебные предметы)
    Cols 2-5: Hours for Grades 1, 2, 3, 4
    Col 6: Form of interim assessment
    """
    results: List[CurriculumRequirement] = []
    seen: Set[Tuple[int, Optional[str], str, float]] = set()

    for r_idx in range(len(table.rows)):
        row = table.rows[r_idx]
        if len(row.cells) < 6:
            continue

        c0 = normalize_text(row.cells[0].text)
        c1 = normalize_text(row.cells[1].text)
        subj = c1 if c1 else c0

        if not subj or is_non_subject_row(subj):
            continue

        for c_idx in range(2, min(6, len(row.cells))):
            grade = c_idx - 1  # Col 2 -> Grade 1, Col 3 -> Grade 2, etc.
            raw_val = row.cells[c_idx].text
            h = parse_hours(raw_val)

            if h is None or h <= 0:
                logger.debug("Skipping grade %d, subject '%s': raw value '%s'", grade, subj, raw_val)
                continue

            key = (grade, None, subj, h)
            if key not in seen:
                seen.add(key)
                results.append(
                    CurriculumRequirement(
                        grade=grade,
                        subject=subj,
                        hours_per_week=h,
                        track=None,
                    )
                )

    return results


def _parse_ooo_table(table: Table) -> List[CurriculumRequirement]:
    """Parse Basic General Education (ООО, Grades 5-9) curriculum table.

    Layout:
    Header rows 1-2 define column mappings for tracks and grades:
    - Row 1: Track / Vertical name (e.g. 'Общеобразовательный', 'Математическая вертикаль', 'ИТ-вертикаль')
    - Row 2: Grade (5 to 9)
    Data rows (mandatory + school-formed part):
    - Col 2: Subject name
    - Cols 3-18: Weekly hours per column
    """
    results: List[CurriculumRequirement] = []
    seen: Set[Tuple[int, Optional[str], str, float]] = set()

    if len(table.rows) < 3:
        return []

    r1_cells = table.rows[1].cells
    r2_cells = table.rows[2].cells
    num_header_cols = min(len(r1_cells), len(r2_cells))

    # Extract column mappings from Header Rows 1 and 2
    col_meta: Dict[int, Tuple[int, Optional[str]]] = {}
    for c_idx in range(num_header_cols):
        track_text = normalize_text(r1_cells[c_idx].text)
        grade_text = normalize_text(r2_cells[c_idx].text)

        m_grade = re.search(r"([5-9])", grade_text)
        if m_grade:
            grade = int(m_grade.group(1))
            # Standard 'Общеобразовательный' or blank -> track is None
            track = None
            if track_text and track_text.lower() != "общеобразовательный" and not is_non_subject_row(track_text):
                track = track_text
            col_meta[c_idx] = (grade, track)

    # Process all rows
    for r_idx in range(3, len(table.rows)):
        row = table.rows[r_idx]
        cells = row.cells
        if len(cells) < 3:
            continue

        c2 = normalize_text(cells[2].text)
        subj = c2
        if not subj or is_non_subject_row(subj):
            continue

        # Check for repeated header rows inside the table (e.g., repeating track banners)
        if any(k in subj.lower() for k in ["вертикаль", "общеобразовательный", "2025-2026"]):
            continue

        for c_idx, (grade, track) in col_meta.items():
            if c_idx >= len(cells):
                continue

            raw_val = cells[c_idx].text
            h = parse_hours(raw_val)
            if h is None or h <= 0:
                logger.debug("Skipping grade %d (track=%s), subject '%s': raw value '%s'", grade, track, subj, raw_val)
                continue

            key = (grade, track, subj, h)
            if key not in seen:
                seen.add(key)
                results.append(
                    CurriculumRequirement(
                        grade=grade,
                        subject=subj,
                        hours_per_week=h,
                        track=track,
                    )
                )

    return results


def _parse_soo_table(table: Table) -> Tuple[List[CurriculumRequirement], Dict[str, str]]:
    """Parse Secondary General Education (СОО, Grades 10-11) curriculum table.

    Layout:
    - Row 1: Profile track names ('Гуманитарный', 'Технологический', 'Социально-экономический', etc.)
    - Row 2: Grade numbers (10, 11) and class letters for Grade 11 ('11 а', '11 б', '11 в')
    - Data rows:
      Col 1 (or Col 0): Subject name
      Cols 2-10: Academic hours for corresponding profile/grade columns
    """
    results: List[CurriculumRequirement] = []
    seen: Set[Tuple[int, Optional[str], str, float]] = set()
    grade11_class_map: Dict[str, str] = {}

    if len(table.rows) < 3:
        return [], {}

    r1_cells = table.rows[1].cells
    r2_cells = table.rows[2].cells
    num_header_cols = min(len(r1_cells), len(r2_cells))

    col_meta: Dict[int, Tuple[int, Optional[str]]] = {}
    for c_idx in range(2, num_header_cols):
        track_text = normalize_text(r1_cells[c_idx].text)
        grade_text = normalize_text(r2_cells[c_idx].text)

        if not grade_text and not track_text:
            continue

        m_class = re.search(r"(\d+)\s*([а-яёa-z])", grade_text, re.IGNORECASE)
        m_grade = re.search(r"(1[01])", grade_text)

        grade = int(m_grade.group(1)) if m_grade else None
        if grade is None:
            continue

        track = track_text if track_text and track_text.lower() != "общеобразовательный" else None

        # Capture Grade 11 class-to-track mapping (e.g. '11а' -> 'Социально-экономический')
        if m_class and track:
            class_key = f"{m_class.group(1)}{m_class.group(2).lower()}"
            grade11_class_map[class_key] = track

        col_meta[c_idx] = (grade, track)

    # Process subject rows
    for r_idx in range(3, len(table.rows)):
        row = table.rows[r_idx]
        cells = row.cells
        if len(cells) < 2:
            continue

        c0 = normalize_text(cells[0].text)
        c1 = normalize_text(cells[1].text)
        subj = c1 if c1 else c0

        if not subj or is_non_subject_row(subj):
            continue

        for c_idx, (grade, track) in col_meta.items():
            if c_idx >= len(cells):
                continue

            raw_val = cells[c_idx].text
            h = parse_hours(raw_val)
            if h is None or h <= 0:
                logger.debug("Skipping grade %d (track=%s), subject '%s': raw value '%s'", grade, track, subj, raw_val)
                continue

            key = (grade, track, subj, h)
            if key not in seen:
                seen.add(key)
                results.append(
                    CurriculumRequirement(
                        grade=grade,
                        subject=subj,
                        hours_per_week=h,
                        track=track,
                    )
                )

    return results, grade11_class_map


def extract_grade11_class_track_mapping(file_path: str) -> Dict[str, str]:
    """Helper to extract class-letter-to-track mapping for Grade 11 from a СОО curriculum plan."""
    doc = docx.Document(file_path)
    if not doc.tables:
        return {}
    _, class_map = _parse_soo_table(doc.tables[0])
    return class_map


def parse_curriculum_docx(file_path: str) -> List[CurriculumRequirement]:
    """Main entry point to parse a school curriculum Word document (.docx) into CurriculumRequirement models.

    Auto-detects the table shape (NOO/OOO/SOO) based on structural heuristics and dispatches
    to the specialized parser.
    """
    doc = docx.Document(file_path)
    if not doc.tables:
        logger.warning("No tables found in %s", file_path)
        return []

    table = doc.tables[0]
    shape = detect_curriculum_shape(table)
    logger.info("Detected curriculum shape '%s' for file: %s", shape, file_path)

    if shape == "NOO":
        return _parse_noo_table(table)
    elif shape == "OOO":
        return _parse_ooo_table(table)
    elif shape == "SOO":
        requirements, _ = _parse_soo_table(table)
        return requirements
    else:
        # Fallback to OOO parser as general multi-column parser
        return _parse_ooo_table(table)
