"""Workload parser for Russian school teacher workload spreadsheets (Нагрузка.xlsx).

Parses teacher-to-class workload matrix from the 'Общая' sheet into
Teacher and Assignment domain models, resolving multi-subject teacher hours
against the curriculum requirements.
"""

import itertools
import logging
import re
from collections import defaultdict
from typing import Dict, List, Optional, Set, Tuple

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

from app.models.schema import Assignment, CurriculumRequirement, Teacher, ValidationIssue

logger = logging.getLogger(__name__)

# Cyrillic to Latin transliteration mapping for stable slug IDs
CYRILLIC_TO_LATIN = {
    "а": "a",
    "б": "b",
    "в": "v",
    "г": "g",
    "д": "d",
    "е": "e",
    "ё": "e",
    "ж": "zh",
    "з": "z",
    "и": "i",
    "й": "y",
    "к": "k",
    "л": "l",
    "м": "m",
    "н": "n",
    "о": "o",
    "п": "p",
    "р": "r",
    "с": "s",
    "т": "t",
    "у": "u",
    "ф": "f",
    "х": "kh",
    "ц": "ts",
    "ч": "ch",
    "ш": "sh",
    "щ": "shch",
    "ъ": "",
    "ы": "y",
    "ь": "",
    "э": "e",
    "ю": "yu",
    "я": "ya",
}


def slugify_name(name: str) -> str:
    """Convert a teacher name or vacancy string to a safe, stable, readable slug ID.

    Examples:
        'Иншакова О.А.' -> 'inshakova-o-a'
        'Вакансия (Трохина Е.А.)' -> 'vakansiya-trokhina-e-a'
        'Вакансия' -> 'vakansiya'
    """
    s = name.lower().strip()
    res: List[str] = []
    for char in s:
        if char in CYRILLIC_TO_LATIN:
            res.append(CYRILLIC_TO_LATIN[char])
        elif char.isalnum():
            res.append(char)
        elif char in " ._-()/,":
            res.append("-")
    slug = "".join(res)
    slug = re.sub(r"-+", "-", slug).strip("-")
    return slug or "unnamed-teacher"


def clean_cell_str(val: Optional[object]) -> str:
    """Normalize cell value to string, removing non-breaking spaces and whitespace."""
    if val is None:
        return ""
    s = str(val).replace("\xa0", " ").strip()
    return re.sub(r"\s+", " ", s).strip()


def parse_cell_hours(val: Optional[object]) -> List[float]:
    """Extract numeric hours per week from a cell value.

    Supports composite expressions like '2+1', '3+1', floats, integers, and comma decimals.
    Returns a list of individual float hours (e.g. '2+1' -> [2.0, 1.0]).
    """
    if val is None:
        return []

    # If openpyxl directly returned numeric value
    if isinstance(val, (int, float)):
        return [float(val)] if val > 0 else []

    cleaned = clean_cell_str(val)
    if not cleaned or cleaned == "-":
        return []

    hours_list: List[float] = []
    for part in cleaned.split("+"):
        p = part.strip().replace(",", ".")
        try:
            h = float(p)
            if h > 0:
                hours_list.append(h)
        except ValueError:
            logger.debug("Could not parse hour component '%s' from cell '%s'", p, val)

    return hours_list


def class_name_to_grade_track(
    class_name: str,
    curriculum: List[CurriculumRequirement],
    soo_class_track_map: Optional[Dict[str, str]] = None,
) -> Tuple[int, Optional[str]]:
    """Resolves a raw class name (e.g. '5А', '11Б') to its (grade, track).

    - Grade is the leading number (e.g. '5А' -> 5, '11Б' -> 11).
    - Track resolution:
      - Grades 1-6: No tracks (track is None).
      - Grades 7-9: In up_ooo.docx, track columns are not assigned fixed class letters in headers;
        track is None here by default (multiple candidate tracks are explored dynamically during hour resolution).
      - Grade 10: No fixed class-letter mapping in up_soo.docx (all tracks apply loosely); track is None.
      - Grade 11: Uses the detected class-letter-to-track mapping (e.g. '11а' -> 'Социально-экономический',
        '11б' -> 'Гуманитарный', '11в' -> 'Технологический...').
    """
    m = re.match(r"^(\d+)", class_name.strip())
    if not m:
        raise ValueError(f"Cannot extract grade level from class name '{class_name}'")

    grade = int(m.group(1))

    if grade <= 6 or grade == 10:
        return grade, None

    if grade == 11 and soo_class_track_map:
        norm_key = class_name.lower().replace(" ", "")
        if norm_key in soo_class_track_map:
            return grade, soo_class_track_map[norm_key]
        for k, v in soo_class_track_map.items():
            if k.lower().replace(" ", "") == norm_key:
                return grade, v

    return grade, None


def normalize_subject_keys_for_grade(subject_label: str, grade: int) -> List[str]:
    """Expands a subject label into standard curriculum subject keys for a specific grade."""
    s = subject_label.lower().strip()
    s = s.replace("обществозна-ние", "обществознание")
    s = s.replace("обзр", "основы безопасности и защиты родины")
    s = s.replace("изо", "изобразительное искусство")

    if s == "технология":
        return ["труд (технология)"]
    if s == "наглядная геометрия":
        return ["наглядная геометрия: конструирование многогранников и тел вращения"]
    if s == "практикум по экономике":
        return ["практикум по экономике", "экономика"]
    if s == "литература" and grade <= 4:
        return ["литературное чтение"]
    if s == "математика":
        if grade >= 10:
            return ["алгебра и начала математического анализа", "геометрия", "вероятность и статистика"]
        elif grade >= 7:
            return ["алгебра", "геометрия", "вероятность и статистика"]
        return ["математика"]

    return [s]


def _format_hours(hours: float) -> str:
    """Formats an hours value for display, dropping a trailing '.0' when whole."""
    if float(hours).is_integer():
        return str(int(hours))
    return f"{hours:g}"


def resolve_multi_subject_hours(
    teacher_id: str,
    teacher_name: str,
    subjects: List[str],
    class_name: str,
    total_hours: float,
    curriculum: List[CurriculumRequirement],
    soo_class_track_map: Optional[Dict[str, str]] = None,
) -> Tuple[List[Assignment], Optional[ValidationIssue]]:
    """Resolves a teacher's total workload hours for a class into subject-specific assignments.

    Cross-references candidate curriculum requirements for (grade, candidate_tracks):
    - If exactly one subject breakdown sums to `total_hours` (within float tolerance 0.01),
      returns individual Assignment objects for each matching subject.
    - If multiple combinations or NO combinations match, falls back to assigning total_hours
      to the first listed subject and generates a ValidationIssue warning.
    """
    if len(subjects) <= 1:
        subj = subjects[0] if subjects else "Не указан"
        return [
            Assignment(
                teacher_id=teacher_id,
                subject=subj,
                class_name=class_name,
                hours_per_week=total_hours,
            )
        ], None

    grade, fixed_track = class_name_to_grade_track(class_name, curriculum, soo_class_track_map)

    # Index curriculum for quick lookup: (grade, track) -> {lower_subject: (exact_name, hours)}
    curric_map: Dict[Tuple[int, Optional[str]], Dict[str, Tuple[str, float]]] = defaultdict(dict)
    tracks_for_grade: Set[Optional[str]] = set()

    for req in curriculum:
        if req.grade == grade:
            tracks_for_grade.add(req.track)
            curric_map[(req.grade, req.track)][req.subject.lower()] = (req.subject, req.hours_per_week)

    # Determine candidate tracks to evaluate
    if fixed_track is not None:
        candidate_tracks = [fixed_track]
    elif grade <= 6:
        candidate_tracks = [None]
    else:
        candidate_tracks = list(tracks_for_grade) if tracks_for_grade else [None]

    valid_breakdowns: List[Tuple[Optional[str], Tuple[Tuple[str, float], ...]]] = []

    for track in candidate_tracks:
        grade_curric = curric_map[(grade, track)]
        available_reqs: List[Tuple[str, float]] = []

        for s in subjects:
            sub_keys = normalize_subject_keys_for_grade(s, grade)
            for sk in sub_keys:
                if sk in grade_curric:
                    req_entry = grade_curric[sk]
                    if req_entry not in available_reqs:
                        available_reqs.append(req_entry)

        # Explore all non-empty combinations of matched curriculum requirements
        for r_len in range(1, len(available_reqs) + 1):
            for combo in itertools.combinations(available_reqs, r_len):
                combo_sum = sum(h for _, h in combo)
                if abs(combo_sum - total_hours) < 0.01:
                    valid_breakdowns.append((track, combo))

    # Deduplicate breakdowns that yield identical subject/hours combinations
    unique_breakdowns: List[Tuple[Tuple[str, float], ...]] = []
    for _, combo in valid_breakdowns:
        sorted_combo = tuple(sorted(combo, key=lambda x: x[0]))
        if sorted_combo not in unique_breakdowns:
            unique_breakdowns.append(sorted_combo)

    if len(unique_breakdowns) == 1:
        breakdown = unique_breakdowns[0]
        assignments = [
            Assignment(
                teacher_id=teacher_id,
                subject=subj_name,
                class_name=class_name,
                hours_per_week=h,
            )
            for subj_name, h in breakdown
        ]
        return assignments, None

    # Ambiguous or unresolved case: fallback to first subject with a warning ValidationIssue
    primary_subject = subjects[0]
    fallback_assignment = [
        Assignment(
            teacher_id=teacher_id,
            subject=primary_subject,
            class_name=class_name,
            hours_per_week=total_hours,
        )
    ]

    hours_str = _format_hours(total_hours)
    subjects_str = ", ".join(subjects)

    if len(unique_breakdowns) > 1:
        mismatch_type = "ambiguous_breakdown"
        message = (
            f"Не удалось точно определить, сколько часов из {hours_str} у преподавателя "
            f"{teacher_name} в классе {class_name} относится к каждому предмету ({subjects_str}) — "
            f"несколько вариантов распределения дают одинаковую сумму. Все часы условно отнесены "
            f"к предмету «{primary_subject}» — проверьте вручную."
        )
    else:
        mismatch_type = "no_valid_breakdown"
        message = (
            f"У преподавателя {teacher_name} в классе {class_name} указано {hours_str} ч./нед. "
            f"по предметам ({subjects_str}), но по учебному плану для этого класса такая сумма часов "
            f"не предусмотрена — возможно, это дополнительный школьный час. Все часы условно отнесены "
            f"к предмету «{primary_subject}» — проверьте вручную."
        )

    issue = ValidationIssue(
        severity="warning",
        message=message,
        context={
            "mismatch_type": mismatch_type,
            "teacher_id": teacher_id,
            "teacher_name": teacher_name,
            "class_name": class_name,
            "grade": grade,
            "total_hours": total_hours,
            "subjects_tried": subjects,
            "candidate_tracks_tried": candidate_tracks,
            "matching_combinations_count": len(unique_breakdowns),
        },
    )

    return fallback_assignment, issue


def parse_workload_xlsx(
    file_path: str,
    curriculum: Optional[List[CurriculumRequirement]] = None,
    soo_class_track_map: Optional[Dict[str, str]] = None,
) -> Tuple[List[Teacher], List[Assignment], List[ValidationIssue]]:
    """Parse the teacher workload spreadsheet (`Нагрузка.xlsx`).

    Reads the 'Общая' sheet (Teacher × Class matrix):
    - Row 1: Class name headers starting at column C (e.g. '1А', '1Б', ... '11В').
    - Column A: Teacher full name (or vacancy name).
    - Column B: Subject label group on the first row of each group, forward-filled to subsequent teacher rows.
    - Column C onwards: Hours per week assigned to the teacher for each class.

    Returns:
        (teachers, assignments, warnings) tuple.
    """
    curriculum = curriculum or []
    soo_class_track_map = soo_class_track_map or {}

    wb = openpyxl.load_workbook(file_path, data_only=True)
    sheet_name = "Общая"
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"Sheet '{sheet_name}' not found in '{file_path}'. Available sheets: {wb.sheetnames}")

    ws: Worksheet = wb[sheet_name]

    # Step 1: Detect class name headers in Row 1 (from column 3 onwards)
    class_columns: List[Tuple[int, str]] = []
    for col_idx in range(3, ws.max_column + 1):
        raw_val = ws.cell(1, col_idx).value
        class_name = clean_cell_str(raw_val)
        if class_name:
            # Match school class patterns like '1А', '5В', '11Б'
            if re.match(r"^\d{1,2}\s*[А-ЯA-Z]$", class_name, re.IGNORECASE):
                # Standardize class name e.g. '1 А' -> '1А'
                norm_class = re.sub(r"\s+", "", class_name).upper()
                class_columns.append((col_idx, norm_class))

    logger.info("Found %d class columns in row 1 of sheet '%s'", len(class_columns), sheet_name)

    # Step 2: Iterate teacher rows (Row 2 to max_row) with subject label forward-fill
    teachers: List[Teacher] = []
    assignments: List[Assignment] = []
    warnings: List[ValidationIssue] = []

    seen_slugs: Dict[str, Teacher] = {}
    slug_counts: Dict[str, int] = {}

    current_subject_label = ""

    for r_idx in range(2, ws.max_row + 1):
        c1_val = clean_cell_str(ws.cell(r_idx, 1).value)
        c2_val = clean_cell_str(ws.cell(r_idx, 2).value)

        # Update forward-fill subject label if a new one is provided
        if c2_val and c2_val.upper() != "ПРЕДМЕТ":
            current_subject_label = c2_val

        if not c1_val:
            # Empty row or spacer row
            continue

        full_name = c1_val

        # Split comma-separated subject labels into a list of subjects
        if current_subject_label:
            raw_subjects = [s.strip() for s in re.split(r"[,;]\s*", current_subject_label) if s.strip()]
        else:
            raw_subjects = []

        # Generate unique and stable slug ID
        base_slug = slugify_name(full_name)
        slug_counts[base_slug] = slug_counts.get(base_slug, 0) + 1
        if slug_counts[base_slug] > 1:
            slug = f"{base_slug}-{slug_counts[base_slug]}"
            logger.warning("Duplicate teacher name '%s' found at row %d, assigned unique slug '%s'", full_name, r_idx, slug)
        else:
            slug = base_slug

        teacher = Teacher(
            id=slug,
            full_name=full_name,
            subjects=raw_subjects,
        )

        if slug in seen_slugs:
            logger.warning("Teacher slug '%s' already exists; skipping duplicate teacher registration", slug)
        else:
            seen_slugs[slug] = teacher
            teachers.append(teacher)

        # Step 3: Extract numeric hours for each class column
        for col_idx, class_name in class_columns:
            cell_val = ws.cell(r_idx, col_idx).value
            hours = parse_cell_hours(cell_val)
            for h in hours:
                resolved_ass, issue = resolve_multi_subject_hours(
                    teacher_id=teacher.id,
                    teacher_name=teacher.full_name,
                    subjects=raw_subjects,
                    class_name=class_name,
                    total_hours=h,
                    curriculum=curriculum,
                    soo_class_track_map=soo_class_track_map,
                )
                assignments.extend(resolved_ass)
                if issue:
                    warnings.append(issue)

    # Step 4: Detect parallel subgroups for (class_name, subject) pairs with multiple teachers
    # When multiple distinct teachers teach the same subject to the same class, assign group="1", "2", ...
    pair_assignments: Dict[Tuple[str, str], List[Assignment]] = defaultdict(list)
    for a in assignments:
        pair_assignments[(a.class_name, a.subject)].append(a)

    for (class_name, subject), a_list in pair_assignments.items():
        distinct_teachers = set(a.teacher_id for a in a_list)
        if len(distinct_teachers) > 1:
            # Assign group numbers in order of teacher occurrence
            teacher_to_group: Dict[str, str] = {}
            curr_group_num = 1
            for a in a_list:
                if a.teacher_id not in teacher_to_group:
                    teacher_to_group[a.teacher_id] = str(curr_group_num)
                    curr_group_num += 1
                a.group = teacher_to_group[a.teacher_id]
        else:
            for a in a_list:
                a.group = None

    logger.info(
        "Parsed %d teachers, %d assignments (%d validation warnings) from sheet '%s'",
        len(teachers),
        len(assignments),
        len(warnings),
        sheet_name,
    )
    return teachers, assignments, warnings
