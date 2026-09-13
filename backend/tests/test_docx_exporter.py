import os
import tempfile
import time
from fastapi.testclient import TestClient

import docx
import pytest

from app.export.docx_exporter import export_schedule_to_docx
from app.main import app
from app.models.schema import Assignment, ScheduleSlot, Teacher
from app.parsers.curriculum_parser import (
    extract_grade11_class_track_mapping,
    parse_curriculum_docx,
)
from app.parsers.workload_parser import parse_workload_xlsx
from app.solver.scheduler import generate_schedule

client = TestClient(app)

FIXTURES_DIR = os.path.join(os.path.dirname(__file__), "fixtures")


def test_synthetic_schedule_docx_export():
    """Test exporting a small synthetic schedule to DOCX."""
    teachers = [
        Teacher(id="ivanov-ii", full_name="Иванов И.И.", subjects=["Математика"]),
        Teacher(id="petrova-aa", full_name="Петрова А.А.", subjects=["Русский язык"]),
        Teacher(id="sidorov-ss", full_name="Сидоров С.С.", subjects=["История", "Обществознание"]),
    ]

    assignments = [
        Assignment(teacher_id="ivanov-ii", subject="Математика", class_name="5А", hours_per_week=2.0),
        Assignment(teacher_id="petrova-aa", subject="Русский язык", class_name="5А", hours_per_week=2.0),
        Assignment(teacher_id="sidorov-ss", subject="История", class_name="5А", hours_per_week=1.0),
        Assignment(teacher_id="sidorov-ss", subject="Обществознание", class_name="5А", hours_per_week=1.0),
    ]

    schedule, status = generate_schedule(assignments, time_limit_seconds=10)
    assert len(schedule) == 6

    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
        output_path = tmp.name

    try:
        export_schedule_to_docx(schedule, teachers, output_path)
        assert os.path.exists(output_path)

        # Open and verify table structure
        doc = docx.Document(output_path)
        assert len(doc.tables) == 1
        table = doc.tables[0]

        # 2 header rows + 4 data rows (Ivanov: Math, Petrova: Rus, Sidorov: History, Sidorov: Social)
        assert len(table.rows) == 6
        assert len(table.columns) == 47  # 2 header cols + 45 period cols

        # Check headers
        assert "Предмет" in table.cell(0, 0).text
        assert "ФИО" in table.cell(0, 1).text
        assert "Понедельник" in table.cell(0, 2).text

        # Verify teacher names appear in Col 1 of data rows
        teacher_names_in_table = [row.cells[1].text.strip() for row in table.rows[2:6]]
        assert "Иванов И.И." in teacher_names_in_table
        assert "Петрова А.А." in teacher_names_in_table
        assert "Сидоров С.С." in teacher_names_in_table

        # Verify that class '5А' is present in scheduled cells
        scheduled_cells = []
        for row in table.rows[2:6]:
            for cell in row.cells[2:47]:
                txt = cell.text.strip()
                if txt:
                    scheduled_cells.append(txt)

        assert len(scheduled_cells) == 6
        assert all(c == "5А" for c in scheduled_cells)
        print(f"\nSynthetic DOCX generated successfully with {len(table.rows)} rows and {len(scheduled_cells)} scheduled lesson cells.")

    finally:
        if os.path.exists(output_path):
            os.remove(output_path)


def test_real_dataset_schedule_docx_export():
    """Test exporting the full 76-class real dataset schedule to DOCX."""
    workload_path = os.path.join(FIXTURES_DIR, "workload.xlsx")
    soo_path = os.path.join(FIXTURES_DIR, "up_soo.docx")

    all_curriculum = (
        parse_curriculum_docx(os.path.join(FIXTURES_DIR, "up_noo.docx"))
        + parse_curriculum_docx(os.path.join(FIXTURES_DIR, "up_ooo.docx"))
        + parse_curriculum_docx(soo_path)
    )
    soo_class_track_map = extract_grade11_class_track_mapping(soo_path)

    teachers, assignments, _ = parse_workload_xlsx(
        file_path=workload_path,
        curriculum=all_curriculum,
        soo_class_track_map=soo_class_track_map,
    )

    schedule, status = generate_schedule(assignments, time_limit_seconds=30)
    print(f"Real dataset solve status: {status.value}, slots: {len(schedule)}")

    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
        output_path = tmp.name

    try:
        t0 = time.time()
        export_schedule_to_docx(schedule, teachers, output_path)
        export_duration = time.time() - t0

        doc = docx.Document(output_path)
        assert len(doc.tables) == 1
        table = doc.tables[0]

        data_rows_count = len(table.rows) - 2
        print(f"\n==================== REAL DATASET DOCX EXPORT ====================")
        print(f"Export Elapsed Time: {export_duration:.2f} seconds")
        print(f"Table Rows: {len(table.rows)} (2 header rows + {data_rows_count} teacher-subject data rows)")
        print(f"Table Columns: {len(table.columns)}")

        # Count total populated lesson cells
        total_lesson_cells = 0
        for row in table.rows[2:]:
            for cell in row.cells[2:]:
                if cell.text.strip():
                    total_lesson_cells += 1

        print(f"Total Scheduled Lesson Cells in Table: {total_lesson_cells}")
        assert len(table.rows) >= 2
        assert len(table.columns) == 47
        assert total_lesson_cells == len(schedule)

    finally:
        if os.path.exists(output_path):
            os.remove(output_path)


def test_export_endpoint_happy_path():
    """Test POST /api/generate-schedule/export returns valid downloadable docx file."""
    up_noo_path = os.path.join(FIXTURES_DIR, "up_noo.docx")
    up_ooo_path = os.path.join(FIXTURES_DIR, "up_ooo.docx")
    up_soo_path = os.path.join(FIXTURES_DIR, "up_soo.docx")
    workload_path = os.path.join(FIXTURES_DIR, "workload.xlsx")

    start_time = time.time()

    with (
        open(up_noo_path, "rb") as f_noo,
        open(up_ooo_path, "rb") as f_ooo,
        open(up_soo_path, "rb") as f_soo,
        open(workload_path, "rb") as f_workload,
    ):
        files = {
            "up_noo": ("up_noo.docx", f_noo, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
            "up_ooo": ("up_ooo.docx", f_ooo, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
            "up_soo": ("up_soo.docx", f_soo, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
            "workload": ("workload.xlsx", f_workload, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
        }

        response = client.post(
            "/api/generate-schedule/export",
            files=files,
            params={"time_limit_seconds": 30},
        )

    elapsed = time.time() - start_time

    print(f"\n==================== EXPORT ENDPOINT API RUN ====================")
    print(f"HTTP Status Code: {response.status_code}")
    print(f"Content-Type: {response.headers.get('content-type')}")
    print(f"Content-Disposition: {response.headers.get('content-disposition')}")
    print(f"Payload Size: {len(response.content)} bytes")
    print(f"Total API Elapsed Time: {elapsed:.2f} seconds")

    assert response.status_code == 200
    assert "application/vnd.openxmlformats-officedocument.wordprocessingml.document" in response.headers.get("content-type", "")
    assert 'attachment; filename="schedule.docx"' in response.headers.get("content-disposition", "")

    # Validate that returned bytes form a valid DOCX document
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
        tmp.write(response.content)
        temp_docx_path = tmp.name

    try:
        downloaded_doc = docx.Document(temp_docx_path)
        assert len(downloaded_doc.tables) == 1
        t = downloaded_doc.tables[0]
        assert len(t.rows) >= 2
        assert len(t.columns) == 47
        print(f"✅ Verified: Downloaded docx is valid with {len(t.rows)} table rows.")
    finally:
        if os.path.exists(temp_docx_path):
            os.remove(temp_docx_path)
